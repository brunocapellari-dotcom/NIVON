"""Leitura estruturada XLSX ANP: série SEMANAL OFICIAL por município.

Layout é contrato EXTERNO que deve ser homologado contra planilhas ANP reais.
Sem configuração documentada não aceita a planilha; não presume cabeçalhos,
não reconstrói semanas e não confunde diesel B de revenda com diesel A/PPI.
"""
from __future__ import annotations

import hashlib
import io
import math
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from .lineup_nacional import instante

MAX_XLSX_BYTES = 24_000_000
MAX_LINHAS = 550_000
UF_VALIDAS=frozenset("AC AL AM AP BA CE DF ES GO MA MG MS MT PA PB PE PI PR RJ RN RO RR RS SC SE SP TO".split())


class ErroANPSemanal(ValueError):
    pass


@dataclass(frozen=True)
class LayoutSemanalANP:
    aba: str
    cabecalho_linha: int
    colunas: dict[str,str]
    documento_schema: str
    unidade_preco: str = "BRL/L"
    # Nomes exatos são capturados do primeiro XLSX real homologado.
    # Layout inventado para teste NÃO pode habilitar catálogo operacional.


OBRIGATORIAS=frozenset({"uf","municipio","produto","semana_inicio",
                        "semana_fim","preco_medio","postos"})


def _layout(layout:LayoutSemanalANP):
    if not isinstance(layout,LayoutSemanalANP):
        raise ErroANPSemanal("layout do XLSX não homologado")
    if not isinstance(layout.aba,str) or not layout.aba.strip() or not isinstance(layout.cabecalho_linha,int) or not 1<=layout.cabecalho_linha<=30:
        raise ErroANPSemanal("aba/cabeçalho não verificados")
    if not isinstance(layout.colunas,dict) or set(layout.colunas)!=OBRIGATORIAS:
        raise ErroANPSemanal("mapa de colunas semanal incompleto")
    if len(set(layout.colunas.values()))!=len(OBRIGATORIAS):
        raise ErroANPSemanal("colunas XLSX mapeadas repetidamente")
    if layout.unidade_preco!="BRL/L" or not isinstance(layout.documento_schema,str) or not re.fullmatch(r"[A-Za-z0-9_:.-]{12,120}",layout.documento_schema):
        raise ErroANPSemanal("unidade/esquema sem evidência documental")


def _data(v, nome):
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    if isinstance(v,str):
        try:return datetime.strptime(v.strip(),"%d/%m/%Y").date()
        except ValueError:pass
    raise ErroANPSemanal(nome+" inválida (data Excel ou DD/MM/AAAA)")


def _decimal(v, nome):
    if isinstance(v,bool) or v is None:
        raise ErroANPSemanal(nome+" não numérico")
    if isinstance(v,str):
        s=v.strip()
        if not re.fullmatch(r"[0-9]{1,3}(?:\.[0-9]{3})*(?:,[0-9]{1,4})?|[0-9]+(?:,[0-9]{1,4})?",s):
            raise ErroANPSemanal(nome+" formato monetário inválido")
        v=s.replace(".","").replace(",",".")
    try:d=Decimal(str(v))
    except (InvalidOperation,TypeError) as exc:
        raise ErroANPSemanal(nome+" não decimal") from exc
    if not d.is_finite():raise ErroANPSemanal(nome+" não finito")
    return d


def interpretar_anp_semanal_xlsx(payload:bytes,*,layout:LayoutSemanalANP,
                                coletado_em:str,
                                municipios_mt:frozenset[str]=frozenset({"SORRISO","SINOP","RONDONÓPOLIS"}),
                                ano_esperado:int|None=None)->dict:
    """Leitura OFFLINE, sem chamadas HTTP, metadados pessoais ou extrapolações.

    Retorna apenas municípios de MT escolhidos e efetivamente pesquisados.
    Qualquer erro na linha de diesel relevante invalida o arquivo inteiro.
    """
    _layout(layout)
    coleta=instante(coletado_em)
    if not isinstance(payload,bytes) or not 100<=len(payload)<=MAX_XLSX_BYTES or not payload.startswith(b"PK\x03\x04"):
        raise ErroANPSemanal("arquivo XLSX ausente/ilegal/muito grande")
    if type(ano_esperado) is not int and ano_esperado is not None:
        raise ErroANPSemanal("ano esperado inválido")
    if not isinstance(municipios_mt,frozenset) or not municipios_mt or any(not isinstance(x,str) or not x for x in municipios_mt):
        raise ErroANPSemanal("municípios devem ser um conjunto explícito")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ErroANPSemanal("instale extra do projeto [excel] para ler planilha ANP") from exc
    try:
        wb=load_workbook(io.BytesIO(payload),read_only=True,data_only=True,keep_links=False)
    except Exception as exc:
        raise ErroANPSemanal("XLSX inválido/não suportado") from exc
    try:
        if layout.aba not in wb.sheetnames:
            raise ErroANPSemanal("aba mudou: homologar novo layout")
        ws=wb[layout.aba]
        linhas=ws.iter_rows(values_only=True)
        header=None
        for i in range(1,layout.cabecalho_linha+1):
            try:header=next(linhas)
            except StopIteration as exc:raise ErroANPSemanal("cabeçalho XLSX não encontrado") from exc
        nomes=[str(x).strip() if x is not None else "" for x in header]
        if len(set(x for x in nomes if x))!=len([x for x in nomes if x]):
            raise ErroANPSemanal("nomes de colunas duplicados")
        if not set(layout.colunas.values()).issubset(set(nomes)):
            raise ErroANPSemanal("cabeçalho não corresponde à amostra homologada")
        ix={k:nomes.index(v) for k,v in layout.colunas.items()}
        saida=[]
        vistos=set()
        encontrados=frozenset()
        linhas_analise=0
        for rownum,row in enumerate(linhas,start=layout.cabecalho_linha+1):
            linhas_analise+=1
            if linhas_analise>MAX_LINHAS:
                raise ErroANPSemanal("planilha excede quantidade segura de linhas")
            if all(x is None or x=="" for x in row):continue
            get=lambda field: row[ix[field]] if ix[field]<len(row) else None
            uf=str(get("uf") or "").strip().upper()
            if uf not in UF_VALIDAS:
                raise ErroANPSemanal(f"linha {rownum}: UF inválida")
            municipio=" ".join(str(get("municipio") or "").strip().upper().split())
            if not municipio:
                raise ErroANPSemanal(f"linha {rownum}: município ausente")
            produto=" ".join(str(get("produto") or "").strip().upper().split())
            if produto not in ("DIESEL","DIESEL S10","ÓLEO DIESEL","ÓLEO DIESEL S10"):
                continue
            if uf!="MT" or municipio not in municipios_mt:
                continue
            di=_data(get("semana_inicio"),"início da semana")
            df=_data(get("semana_fim"),"fim da semana")
            if df<di or (df-di).days!=6:
                raise ErroANPSemanal("semana ANP não tem 7 dias")
            if di>coleta.date() or df.year<2000:
                raise ErroANPSemanal("período no futuro/incompatível")
            if ano_esperado is not None and not (di.year==ano_esperado or df.year==ano_esperado):
                raise ErroANPSemanal("semana fora do ano declarado")
            preco=_decimal(get("preco_medio"),"preço médio de revenda")
            if not Decimal("0.1") <= preco <= Decimal("50"):
                raise ErroANPSemanal("preço médio de diesel fora de faixa")
            postos=get("postos")
            if type(postos) is not int or not 1<=postos<=15000:
                raise ErroANPSemanal("número de postos não inteiro/positivo")
            chave=(uf,municipio,produto,di.isoformat(),df.isoformat())
            if chave in vistos:
                raise ErroANPSemanal("linha da mesma semana/município/produto duplicada")
            vistos.add(chave)
            saida.append({
                "uf":uf,"municipio":municipio,"produto":produto,
                "semana_inicio_oficial":di.isoformat(),
                "semana_fim_oficial":df.isoformat(),
                "preco_medio_revenda_rs_l":float(preco),
                "postos_pesquisados":postos,
                "natureza":"ANP_SEMANA_OFICIAL_MUNICIPIO_MEDIA_ARITMETICA_REVENDA",
            })
        if not saida:
            raise ErroANPSemanal("nenhum município/produto alvo encontrado — não concluir que preço é zero")
        return {"dataset_id":"ANP_PRECO_REVENDA_DIESEL",
                "fonte":"ANP_SEMANAL_OFICIAL_MUNICIPIO",
                "origem":"XLSX_SEMANAL_NAO_CSV_MENSAL",
                "unidade":"BRL/L",
                "coletado_em":coleta.isoformat(),
                "arquivo_sha256":hashlib.sha256(payload).hexdigest(),
                "documento_layout":layout.documento_schema,
                "linhas_lidas":linhas_analise,
                "observacoes":sorted(saida,key=lambda x:(x["semana_fim_oficial"],x["municipio"],x["produto"])),
                "cobertura_municipal":[x for x in sorted(municipios_mt)
                                         if any(y["municipio"]==x for y in saida)],
                "publicacao_autorizada":False,
                "nao_inferir":"DIESEL_A_REFINARIA_PPI_CUSTO_TRANSPORTE_PRECO_ATUAL_SEM_DATA",
        }
    finally:
        wb.close()
