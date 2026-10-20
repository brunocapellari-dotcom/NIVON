"""Parser offline do resumo SEMANAL oficial ANP/LPC, aba MUNICIPIOS.

A planilha consolidada semanal difere do CSV mensal de POSTOS e do histórico
anual de municípios. Layout indicado por documentação pública e ETLs terceiros;
não declarado homologado sem amostras XLSX reais e verificação de direitos.
Sem chamadas de rede, dados por posto ou inferências de semanas.
"""
from __future__ import annotations

import hashlib
import io
import math
import re
import unicodedata
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile, ZipFile

from .lineup_nacional import instante


class ErroResumoANP(ValueError):
    pass


URL_INDICE = ("https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/"
              "precos/levantamento-de-precos-de-combustiveis-ultimas-semanas-pesquisadas")
UF_ESTADOS = {
    "MATO GROSSO":"MT", "MATO GROSSO DO SUL":"MS", "GOIAS":"GO", "DISTRITO FEDERAL":"DF",
    "ACRE":"AC", "ALAGOAS":"AL", "AMAPA":"AP", "AMAZONAS":"AM", "BAHIA":"BA",
    "CEARA":"CE", "ESPIRITO SANTO":"ES", "MARANHAO":"MA", "MINAS GERAIS":"MG",
    "PARA":"PA", "PARAIBA":"PB", "PARANA":"PR", "PERNAMBUCO":"PE", "PIAUI":"PI",
    "RIO DE JANEIRO":"RJ", "RIO GRANDE DO NORTE":"RN", "RIO GRANDE DO SUL":"RS",
    "RONDONIA":"RO", "RORAIMA":"RR", "SANTA CATARINA":"SC", "SAO PAULO":"SP",
    "SERGIPE":"SE", "TOCANTINS":"TO",
}
HEADERS = {
    "inicio": "DATA INICIAL", "fim":"DATA FINAL", "estado":"ESTADO",
    "municipio":"MUNICIPIO", "produto":"PRODUTO",
    "postos":"NUMERO DE POSTOS PESQUISADOS", "unidade":"UNIDADE DE MEDIDA",
    "preco":"PRECO MEDIO REVENDA",
}
MAX_BYTES=14_000_000
MAX_ROWS=90_000


def _canon(value: object) -> str:
    s=unicodedata.normalize("NFKD",str(value or "").upper())
    return " ".join("".join(c for c in s if not unicodedata.combining(c)).split())


def _date(value: object) -> date:
    if isinstance(value,datetime):
        return value.date()
    if isinstance(value,date):
        return value
    if isinstance(value,(int,float)) and not isinstance(value,bool):
        if not math.isfinite(value) or value<40000 or value>90000:
            raise ErroResumoANP("serial Excel fora de intervalo aceitável")
        return date(1899,12,30)+timedelta(days=int(value))
    if isinstance(value,str):
        for fmt in ("%d/%m/%Y","%Y-%m-%d"):
            try:
                return datetime.strptime(value.strip(),fmt).date()
            except ValueError:
                pass
    raise ErroResumoANP("data semanal não reconhecida")


def _decimal(value: object) -> Decimal:
    if isinstance(value,bool) or value is None:
        raise ErroResumoANP("preço não numérico")
    if isinstance(value,str):
        val=value.strip()
        if not re.fullmatch(r"[0-9]+(?:\.[0-9]{3})*(?:,[0-9]{1,4})?|[0-9]+(?:,[0-9]{1,4})?",val):
            raise ErroResumoANP("preço com formato desconhecido")
        value=val.replace(".","").replace(",",".")
    try:
        dec=Decimal(str(value))
    except (ValueError,TypeError,InvalidOperation) as exc:
        raise ErroResumoANP("preço inválido") from exc
    if not dec.is_finite() or not Decimal("0.1")<=dec<=Decimal("50"):
        raise ErroResumoANP("preço de revenda fora dos limites técnicos")
    return dec


def _produto(text: object) -> str | None:
    t=_canon(text)
    if t in {"OLEO DIESEL S10","DIESEL S10"}:
        return "DIESEL S10"
    if t in {"OLEO DIESEL","DIESEL"}:
        return "DIESEL"
    return None


def _verificar_zip(payload: bytes):
    if not isinstance(payload,bytes) or not 100<=len(payload)<=MAX_BYTES or not payload.startswith(b"PK\x03\x04"):
        raise ErroResumoANP("XLSX ausente, não ZIP ou acima do limite")
    try:
        with ZipFile(io.BytesIO(payload)) as z:
            files=z.infolist()
            if (len(files)>400 or sum(f.file_size for f in files)>95_000_000 or
                any(f.flag_bits&1 or f.filename.lower().endswith(("vbaproject.bin",".exe",".dll"))
                    for f in files)):
                raise ErroResumoANP("arquivo Excel expandido/executável/criptografado não permitido")
    except BadZipFile as exc:
        raise ErroResumoANP("XLSX corrompido") from exc


def ler_resumo_semanal_lpc(payload: bytes, *, coletado_em: str,
                          origem_url: str, municipios_mt: frozenset[str] | None = None,
                          aba: str = "MUNICIPIOS",
                          cabecalho_linha: int = 7) -> dict:
    """Lê somente dados MT de diesel. Não publica e não presume semana atual.

    Defaults do cabeçalho são CANDIDATOS, sujeitos a comparação com XLSX oficial.
    """
    captura=instante(coletado_em)
    if not (isinstance(origem_url,str) and origem_url.startswith("https://www.gov.br/anp/")
            and "/precos/arquivos-lpc/" in origem_url and origem_url.endswith(".xlsx")):
        raise ErroResumoANP("URL não corresponde a resumo semanal ANP identificado")
    if not isinstance(aba,str) or aba!="MUNICIPIOS" or type(cabecalho_linha) is not int or not 1<=cabecalho_linha<=30:
        raise ErroResumoANP("layout de resumo semanal requer aba MUNICIPIOS e linha explícita")
    if municipios_mt is not None:
        if not isinstance(municipios_mt,frozenset) or not municipios_mt:
            raise ErroResumoANP("municípios solicitados inválidos")
        wanted={_canon(s) for s in municipios_mt}
        if any(not s for s in wanted):
            raise ErroResumoANP("município de MT em branco")
    else:
        wanted=None
    _verificar_zip(payload)
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ErroResumoANP("dependência opcional nivon[excel] necessária") from exc
    try:
        wb=load_workbook(io.BytesIO(payload),read_only=True,data_only=True,keep_links=False)
    except Exception as exc:
        raise ErroResumoANP("não foi possível ler a planilha XLSX") from exc
    try:
        if aba not in wb.sheetnames:
            raise ErroResumoANP("aba MUNICIPIOS não encontrada; schema mudou")
        rows=wb[aba].iter_rows(values_only=True)
        header=None
        for _ in range(cabecalho_linha):
            try:
                header=next(rows)
            except StopIteration as exc:
                raise ErroResumoANP("cabeçalho ausente") from exc
        cols=[_canon(x) for x in header]
        if len([c for c in cols if c])!=len(set(c for c in cols if c)):
            raise ErroResumoANP("colunas duplicadas no XLSX")
        if not set(HEADERS.values()).issubset(cols):
            raise ErroResumoANP("colunas divergentes do layout semanal LPC documentado")
        ix={k:cols.index(v) for k,v in HEADERS.items()}
        series=[];keys=set();total=0
        for n,row in enumerate(rows,start=cabecalho_linha+1):
            total+=1
            if total>MAX_ROWS:
                raise ErroResumoANP("limite seguro de linhas excedido")
            if not row or all(v is None or v=="" for v in row):
                continue
            def val(k):
                return row[ix[k]] if ix[k]<len(row) else None
            if _canon(val("estado")) not in UF_ESTADOS:
                # Algumas edições incluem rodapés; sem inventar observações.
                if all(val(k) in ("",None) for k in ("inicio","fim","municipio","produto")):
                    continue
                raise ErroResumoANP(f"linha {n}: estado não identificado")
            if UF_ESTADOS[_canon(val("estado"))]!="MT":
                continue
            cidade=_canon(val("municipio"))
            if not cidade:
                raise ErroResumoANP(f"linha {n}: município ausente")
            if wanted is not None and cidade not in wanted:
                continue
            prod=_produto(val("produto"))
            if prod is None:
                continue
            inicio,fim=_date(val("inicio")),_date(val("fim"))
            if (fim-inicio).days!=6 or fim>captura.date() or fim.year<2013:
                raise ErroResumoANP(f"linha {n}: período semanal inválido ou futuro")
            unidade=_canon(val("unidade")).replace(" ","")
            if unidade not in ("R$/L","R$/LITRO"):
                raise ErroResumoANP(f"linha {n}: preço não expresso em R$/litro")
            postos=val("postos")
            if isinstance(postos,bool) or not isinstance(postos,(int,float)) or int(postos)!=postos or not 1<=postos<=15000:
                raise ErroResumoANP(f"linha {n}: número de postos inválido")
            preco=_decimal(val("preco"))
            key=(cidade,prod,inicio.isoformat(),fim.isoformat())
            if key in keys:
                raise ErroResumoANP("observação municipal semanal duplicada")
            keys.add(key)
            series.append({
                "uf":"MT","municipio":cidade,"produto":prod,
                "semana_inicio_oficial":inicio.isoformat(),
                "semana_fim_oficial":fim.isoformat(),
                "preco_medio_revenda_rs_l":float(preco),
                "postos_pesquisados":int(postos),
                "natureza":"MEDIA_AMOSTRAL_REVENDA_MUNICIPIO_ANP_SEMANA_OFICIAL",
                "referencia_fonte":"RESUMO_SEMANAL_LPC_MUNICIPIOS",
            })
        if not series:
            raise ErroResumoANP("arquivo sem município/produto MT solicitado; não afirmar zero")
        return {
            "dataset_id":"ANP_PRECO_REVENDA_DIESEL",
            "fonte":"ANP_RESUMO_LPC_SEMANAL",
            "sha256":hashlib.sha256(payload).hexdigest(),
            "url_origem":origem_url,"coletado_em":captura.isoformat(),
            "linha_cabecalho":cabecalho_linha,"aba":aba,
            "observacoes":sorted(series,key=lambda x:(x["semana_fim_oficial"],x["municipio"],x["produto"])),
            "municipios_distintos":len({x["municipio"] for x in series}),
            "semanas_distintas":len({x["semana_fim_oficial"] for x in series}),
            "publicavel":False,
            "observacao":"Preço municipal é média simples amostral; não é diesel A, PPI, preço de frete ou negociação.",
        }
    finally:
        wb.close()
