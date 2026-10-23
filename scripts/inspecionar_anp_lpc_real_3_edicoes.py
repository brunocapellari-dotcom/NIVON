"""Inspeção controlada de 3 XLSX reais ANP/LPC em runner efêmero.

Não imprime preços, postos individualizados ou dados por estabelecimento.
Não distribui RAW nem cria dados no GitHub/Google Cloud.
Só o índice oficial fornece URLs, sem contornar mecanismos de proteção.
Saída JSON estruturada para evidências de layout e status de ingestão.
"""
from __future__ import annotations

import io
import json
import sqlite3
import tempfile
from pathlib import Path

from nivon.anp_indice_lpc import URL_INDICE, descobrir_semanas_lpc
from nivon.anp_lpc_coletor_real import obter_documento, coletar_lpc
from nivon.anp_lpc_staging_real import (
    painel_qualidade_lpc, auditar_arquivos_origem_lpc,
)
from nivon.anp_resumo_lpc import HEADERS, _canon
from datetime import datetime, timezone


def _layout(xlsx: bytes) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(xlsx), read_only=True,
                       data_only=True, keep_links=False)
    try:
        sheets = []
        for ws in wb:
            ident = []
            for i, row in enumerate(ws.iter_rows(min_row=1, max_row=20,
                                                  values_only=True), start=1):
                titles = {_canon(x) for x in row if isinstance(x, str)}
                matches = len(titles.intersection(HEADERS.values()))
                if matches >= 5:
                    ident.append({"linha": i, "colunas_reconhecidas": matches})
            sheets.append({
                "aba": ws.title, "linhas_declaradas": ws.max_row,
                "colunas_declaradas": ws.max_column,
                "possiveis_cabecalhos": ident[:3],
            })
        return {"abas": sheets, "possui_municipios": "MUNICIPIOS" in wb.sheetnames}
    finally:
        wb.close()


def executar() -> int:
    metadata = {
        "fonte": "ANP_LPC_PUBLICA",
        "escopo": "INSPECAO_TEMPORARIA_NAO_COMERCIAL_SEM_PRECOS",
        "publicavel": False, "dados_enviados_github": False,
        "dados_enviados_google_cloud": False,
    }
    try:
        indice = obter_documento(URL_INDICE, tipo="HTML")
        publicados = descobrir_semanas_lpc(indice, ano=2026)
        semanas = sorted(
            publicados["semanas"], key=lambda x: x["semana_fim"],
            reverse=True)[:3]
        if len(semanas) != 3:
            raise ValueError("índice não exibe três semanas oficiais válidas")
        arquivos, esquemas = {}, []
        for item in semanas:
            url = item["url_pagina"]
            dados = obter_documento(url, tipo="XLSX")
            arquivos[url] = dados
            esquemas.append({
                "semana_inicio": item["semana_inicio"],
                "semana_fim": item["semana_fim"],
                "bytes_lidos": len(dados),
                "layout": _layout(dados),
            })
        with tempfile.TemporaryDirectory(prefix="nivon_anp_real_") as raiz:
            base=Path(raiz)
            with sqlite3.connect(base/"staging.sqlite") as db:
                recebido = coletar_lpc(
                    db, ano=2026, limite_semanas=3,
                    indice_bytes=indice, arquivos_locais=arquivos,
                    raw_dir=base, esperar_segundos=0,
                )
                corte = datetime.now(timezone.utc).isoformat()
                qualidade = painel_qualidade_lpc(db, as_of=corte)
                integridade = auditar_arquivos_origem_lpc(
                    db, raiz_raw=str(base), as_of=corte,
                )
        metadata.update({
            "status": "INSPECAO_COMPLETA" if recebido["status"]=="LOTE_COMPLETO"
                      and integridade["status"]=="RAW_VERIFICADO_LOCALMENTE"
                      else "INSPECAO_COM_PENDENCIAS",
            "planilhas_baixadas": len(arquivos),
            "semanas": esquemas,
            "recibos": [{
                "semana_fim": x["semana_fim"],
                "observacoes_mt": x["observacoes"],
                "municipios_mt": x["municipios"],
                "sha256": x["sha256"],
            } for x in recebido["edicoes"]],
            "falhas": recebido["erros"],
            "qualidade": qualidade,
            "integridade": {
                "status": integridade["status"],
                "quantidade_falhas": integridade["quantidade_falhas"],
            },
            "raw_persistido": False,
        })
        print(json.dumps(metadata, ensure_ascii=False, indent=2))
        return 0 if metadata["status"]=="INSPECAO_COMPLETA" else 2
    except (OSError, ValueError, PermissionError, ImportError,
            sqlite3.Error) as exc:
        metadata.update({
            "status": "INSPECAO_NAO_CONCLUIDA",
            "erro_tipo": type(exc).__name__, "motivo": str(exc)[:280],
            "precos_publicados": False,
        })
        print(json.dumps(metadata, ensure_ascii=False, indent=2))
        return 2


if __name__ == "__main__":
    raise SystemExit(executar())
